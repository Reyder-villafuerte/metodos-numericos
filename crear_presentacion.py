from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.dml import MSO_THEME_COLOR
from pptx.util import Inches, Pt


OUT = "Presentacion_Metodos_No_Lineales.pptx"
W, H = 13.333, 7.5
BG = RGBColor(8, 12, 26)
PANEL = RGBColor(17, 24, 43)
PURPLE = RGBColor(139, 92, 246)
CYAN = RGBColor(45, 212, 191)
WHITE = RGBColor(241, 245, 249)
MUTED = RGBColor(168, 180, 201)
ORANGE = RGBColor(251, 146, 60)


def box(slide, x, y, w, h, color=PANEL, radius=False, line=None):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
        Inches(x), Inches(y), Inches(w), Inches(h),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.color.rgb = line or color
    if radius:
        shape.adjustments[0] = 0.08
    return shape


def text(slide, value, x, y, w, h, size=20, color=WHITE, bold=False,
         align=PP_ALIGN.LEFT, font="Aptos", margin=0.08):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(margin)
    tf.margin_right = Inches(margin)
    tf.margin_top = Inches(margin)
    tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = value
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return shape


def rich_lines(slide, lines, x, y, w, h, size=18, color=WHITE, gap=5):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0.12)
    tf.margin_right = Inches(0.08)
    tf.margin_top = Inches(0.08)
    for i, (label, body, accent) in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(gap)
        p.level = 0
        r = p.add_run()
        r.text = label
        r.font.name = "Aptos"
        r.font.size = Pt(size)
        r.font.bold = True
        r.font.color.rgb = accent
        r2 = p.add_run()
        r2.text = body
        r2.font.name = "Aptos"
        r2.font.size = Pt(size)
        r2.font.color.rgb = color
    return shape


def base(prs, title, section, number):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = BG
    box(slide, 0, 0, W, 0.08, PURPLE)
    text(slide, section.upper(), 0.55, 0.25, 4.5, 0.28, 10, CYAN, True)
    text(slide, title, 0.55, 0.58, 12.0, 0.62, 28, WHITE, True)
    text(slide, f"{number:02d}  |  Métodos numéricos", 10.55, 7.08, 2.25, 0.22, 9, MUTED, False, PP_ALIGN.RIGHT)
    return slide


def bullet(slide, value, x, y, w, color=WHITE, size=18):
    text(slide, "•", x, y, 0.28, 0.35, size, PURPLE, True)
    text(slide, value, x + 0.32, y, w - 0.32, 0.42, size, color)


prs = Presentation()
prs.slide_width = Inches(W)
prs.slide_height = Inches(H)

# 1
s = prs.slides.add_slide(prs.slide_layouts[6])
s.background.fill.solid()
s.background.fill.fore_color.rgb = BG
box(s, 0, 0, W, 0.12, PURPLE)
box(s, 0.68, 1.05, 0.14, 4.95, CYAN)
text(s, "MÉTODOS NUMÉRICOS", 1.15, 1.12, 6.2, 0.35, 13, CYAN, True)
text(s, "Sistemas no lineales", 1.1, 1.63, 10.8, 0.72, 35, WHITE, True)
text(s, "Newton-Raphson para sistemas\n& Gauss-Seidel no lineal", 1.1, 2.45, 10.8, 1.25, 29, PURPLE, True)
text(s, "Presentación académica  |  Unidad 2", 1.12, 4.18, 6.5, 0.4, 17, MUTED)
box(s, 1.1, 5.2, 4.6, 0.72, PANEL, True)
text(s, "Expositores: ____________________", 1.28, 5.32, 4.2, 0.42, 15, WHITE)
text(s, "Objetivo: resolver F(x) = 0 de forma iterativa", 1.1, 6.35, 8.0, 0.35, 14, MUTED)
text(s, "01", 11.6, 5.55, 1.0, 0.75, 36, CYAN, True, PP_ALIGN.CENTER)

# 2
s = base(prs, "¿Qué problema resolvemos?", "Contexto", 2)
text(s, "Muchos modelos de ingeniería, ciencia y economía no pueden despejarse con una fórmula cerrada.", 0.75, 1.45, 11.8, 0.6, 21, WHITE, True)
box(s, 0.75, 2.35, 5.65, 3.45, PANEL, True)
text(s, "Sistema no lineal", 1.05, 2.63, 4.9, 0.4, 22, CYAN, True)
text(s, "F₁(x₁, x₂, …, xₙ) = 0\nF₂(x₁, x₂, …, xₙ) = 0\n⋮\nFₙ(x₁, x₂, …, xₙ) = 0", 1.05, 3.25, 4.9, 1.9, 23, WHITE, True, PP_ALIGN.CENTER)
box(s, 6.75, 2.35, 5.8, 3.45, PANEL, True)
text(s, "Meta numérica", 7.05, 2.63, 4.9, 0.4, 22, PURPLE, True)
bullet(s, "Encontrar x* tal que F(x*) ≈ 0", 7.05, 3.22, 5.0)
bullet(s, "Partir de una aproximación inicial x⁽⁰⁾", 7.05, 3.88, 5.0)
bullet(s, "Controlar error, residuo y convergencia", 7.05, 4.54, 5.0)
text(s, "La calidad de x⁽⁰⁾ influye directamente en el resultado.", 1.0, 6.25, 11.4, 0.4, 17, ORANGE, True, PP_ALIGN.CENTER)

# 3
s = base(prs, "Ejemplo común para comparar", "Modelo", 3)
text(s, "Usaremos el mismo sistema en ambos métodos:", 0.8, 1.35, 6.5, 0.45, 20, WHITE, True)
box(s, 0.8, 2.05, 6.15, 2.65, PANEL, True)
text(s, "F₁(x,y) = x² + y − 3.75 = 0\nF₂(x,y) = x + y² − 3.75 = 0", 1.12, 2.52, 5.5, 1.45, 26, CYAN, True, PP_ALIGN.CENTER)
box(s, 7.55, 2.05, 4.95, 2.65, PANEL, True)
text(s, "Solución conocida", 7.9, 2.42, 4.2, 0.35, 20, PURPLE, True)
text(s, "x* = 1.5\n y* = 1.5", 8.2, 3.03, 3.6, 0.95, 30, WHITE, True, PP_ALIGN.CENTER)
text(s, "Aproximación inicial para la demostración:  x⁽⁰⁾ = (1, 1)", 1.0, 5.25, 11.3, 0.45, 19, ORANGE, True, PP_ALIGN.CENTER)
text(s, "La solución se verifica sustituyendo en las dos ecuaciones: F₁(1.5,1.5)=F₂(1.5,1.5)=0.", 1.0, 6.05, 11.3, 0.45, 15, MUTED, False, PP_ALIGN.CENTER)

# 4
s = base(prs, "Newton-Raphson para sistemas", "Método 1", 4)
text(s, "Idea: reemplazar el sistema no lineal por una aproximación lineal local.", 0.8, 1.32, 11.4, 0.45, 20, WHITE, True)
box(s, 0.8, 2.0, 5.8, 3.6, PANEL, True)
text(s, "Paso fundamental", 1.12, 2.32, 5.1, 0.35, 20, CYAN, True)
text(s, "J(x⁽ᵏ⁾) · Δx = −F(x⁽ᵏ⁾)\n\nx⁽ᵏ⁺¹⁾ = x⁽ᵏ⁾ + Δx", 1.12, 2.95, 5.05, 1.75, 25, WHITE, True, PP_ALIGN.CENTER)
box(s, 7.05, 2.0, 5.45, 3.6, PANEL, True)
text(s, "Jacobiano", 7.4, 2.32, 4.7, 0.35, 20, PURPLE, True)
text(s, "J(x,y) = [ ∂F₁/∂x   ∂F₁/∂y ]\n       [ ∂F₂/∂x   ∂F₂/∂y ]", 7.48, 3.0, 4.55, 1.25, 21, WHITE, True, PP_ALIGN.CENTER)
text(s, "En el proyecto se admite Jacobiano analítico o aproximado por diferencias centradas.", 1.0, 6.1, 11.3, 0.4, 15, MUTED, False, PP_ALIGN.CENTER)

# 5
s = base(prs, "Newton: algoritmo y criterios", "Método 1", 5)
box(s, 0.8, 1.35, 5.55, 4.95, PANEL, True)
text(s, "Algoritmo", 1.12, 1.68, 4.8, 0.35, 21, CYAN, True)
steps = [
    ("1.", "Elegir x⁽⁰⁾ y tolerancia."),
    ("2.", "Evaluar F(x⁽ᵏ⁾) y J(x⁽ᵏ⁾)."),
    ("3.", "Resolver J·Δx = −F."),
    ("4.", "Actualizar x⁽ᵏ⁺¹⁾ = x⁽ᵏ⁾ + Δx."),
    ("5.", "Repetir hasta converger."),
]
for i, (n, val) in enumerate(steps):
    text(s, n, 1.15, 2.25 + i * 0.68, 0.35, 0.35, 18, PURPLE, True)
    text(s, val, 1.58, 2.25 + i * 0.68, 4.35, 0.4, 17, WHITE)
box(s, 6.75, 1.35, 5.75, 4.95, PANEL, True)
text(s, "¿Cuándo termina?", 7.08, 1.68, 4.9, 0.35, 21, PURPLE, True)
bullet(s, "||Δx||∞ ≤ tolerancia", 7.08, 2.35, 4.8)
bullet(s, "||F(x⁽ᵏ⁺¹⁾)||∞ ≤ tolerancia", 7.08, 3.05, 4.8)
bullet(s, "Jacobiano no singular", 7.08, 3.75, 4.8)
text(s, "Ventaja: convergencia muy rápida cerca de la solución.\nRiesgo: puede fallar con un mal punto inicial.", 7.08, 4.65, 4.95, 0.95, 16, ORANGE, True)

# 6
s = base(prs, "Ejemplo: iteraciones de Newton", "Método 1", 6)
text(s, "Para el sistema de la diapositiva 3, con x⁽⁰⁾=(1,1):", 0.8, 1.3, 8.5, 0.4, 19, WHITE, True)
headers = ["k", "x₁", "x₂", "||Δx||∞", "||F(x)||∞"]
rows = [
    ["0", "1.000000", "1.000000", "—", "1.750e+00"],
    ["1", "1.583333", "1.583333", "5.833e-01", "3.403e-01"],
    ["2", "1.501667", "1.501667", "8.167e-02", "6.669e-03"],
    ["3", "1.500001", "1.500001", "1.666e-03", "2.775e-06"],
    ["4", "1.500000", "1.500000", "6.939e-07", "4.814e-13"],
]
x0, y0, row_h = 0.85, 2.0, 0.62
col_w = [0.75, 2.0, 2.0, 2.15, 2.15]
for j, h in enumerate(headers):
    box(s, x0 + sum(col_w[:j]), y0, col_w[j], row_h, PURPLE, False)
    text(s, h, x0 + sum(col_w[:j]) + 0.03, y0 + 0.05, col_w[j] - 0.06, row_h - 0.1, 14, WHITE, True, PP_ALIGN.CENTER)
for i, row in enumerate(rows):
    for j, value in enumerate(row):
        color = CYAN if i == len(rows) - 1 else WHITE
        box(s, x0 + sum(col_w[:j]), y0 + row_h * (i + 1), col_w[j], row_h, PANEL, False, RGBColor(45, 55, 78))
        text(s, value, x0 + sum(col_w[:j]) + 0.03, y0 + row_h * (i + 1) + 0.05, col_w[j] - 0.06, row_h - 0.1, 14, color, i == len(rows) - 1, PP_ALIGN.CENTER)
box(s, 10.15, 2.0, 2.35, 3.72, PANEL, True)
text(s, "Resultado", 10.45, 2.35, 1.75, 0.35, 20, CYAN, True, PP_ALIGN.CENTER)
text(s, "x ≈ 1.5\ny ≈ 1.5", 10.45, 3.05, 1.75, 0.85, 27, WHITE, True, PP_ALIGN.CENTER)
text(s, "Convergió en 4 iteraciones", 10.35, 4.48, 1.95, 0.55, 15, PURPLE, True, PP_ALIGN.CENTER)
text(s, "La reducción acelerada del residuo muestra la precisión alcanzada.", 0.95, 6.25, 11.3, 0.35, 15, MUTED, False, PP_ALIGN.CENTER)

# 7
s = base(prs, "Gauss-Seidel no lineal", "Método 2", 7)
text(s, "Convierte el sistema a una forma de punto fijo:", 0.8, 1.35, 8.0, 0.42, 20, WHITE, True)
box(s, 0.8, 2.0, 5.95, 3.65, PANEL, True)
text(s, "x = g₁(x,y) = √(3.75 − y)\ny = g₂(x,y) = √(3.75 − x)", 1.15, 2.68, 5.25, 1.5, 24, CYAN, True, PP_ALIGN.CENTER)
box(s, 7.15, 2.0, 5.35, 3.65, PANEL, True)
text(s, "Actualización inmediata", 7.5, 2.35, 4.65, 0.35, 20, PURPLE, True)
text(s, "x⁽ᵏ⁺¹⁾ = g₁(x⁽ᵏ⁾, y⁽ᵏ⁾)\ny⁽ᵏ⁺¹⁾ = g₂(x⁽ᵏ⁺¹⁾, y⁽ᵏ⁾)", 7.5, 3.02, 4.65, 1.45, 22, WHITE, True, PP_ALIGN.CENTER)
text(s, "La segunda ecuación usa inmediatamente el nuevo valor de x.", 0.95, 6.12, 11.3, 0.4, 16, ORANGE, True, PP_ALIGN.CENTER)

# 8
s = base(prs, "Gauss-Seidel: algoritmo y convergencia", "Método 2", 8)
box(s, 0.8, 1.35, 6.05, 4.95, PANEL, True)
text(s, "Proceso", 1.12, 1.7, 4.8, 0.35, 21, CYAN, True)
gs_steps = [
    ("1.", "Elegir x⁽⁰⁾ y tolerancia."),
    ("2.", "Despejar cada variable como gᵢ."),
    ("3.", "Actualizar x₁, luego x₂, …, xₙ."),
    ("4.", "Calcular ||x⁽ᵏ⁺¹⁾−x⁽ᵏ⁾||∞."),
    ("5.", "Parar si el error es pequeño."),
]
for i, (n, val) in enumerate(gs_steps):
    text(s, n, 1.15, 2.25 + i * 0.68, 0.35, 0.35, 18, PURPLE, True)
    text(s, val, 1.58, 2.25 + i * 0.68, 4.9, 0.4, 17, WHITE)
box(s, 7.25, 1.35, 5.25, 4.95, PANEL, True)
text(s, "Condiciones prácticas", 7.58, 1.7, 4.5, 0.35, 21, PURPLE, True)
bullet(s, "La forma g debe favorecer contracción.", 7.58, 2.38, 4.45)
bullet(s, "El orden de actualización importa.", 7.58, 3.08, 4.45)
bullet(s, "Es simple, pero suele ser más lento.", 7.58, 3.78, 4.45)
text(s, "En el proyecto, el criterio es:\n||x⁽ᵏ⁺¹⁾ − x⁽ᵏ⁾||∞ ≤ tolerancia", 7.58, 4.65, 4.5, 0.8, 17, ORANGE, True, PP_ALIGN.CENTER)

# 9
s = base(prs, "Comparación de resultados", "Análisis", 9)
headers = ["Aspecto", "Newton-Raphson", "Gauss-Seidel no lineal"]
data = [
    ("Idea", "Linealización local", "Punto fijo con actualización inmediata"),
    ("Información", "Jacobiano", "Despejes gᵢ"),
    ("Velocidad", "Muy rápida cerca de la raíz", "Generalmente más gradual"),
    ("Riesgo", "Jacobiano singular / mala semilla", "No convergencia de la forma g"),
    ("Ejemplo", "4 iteraciones", "Más iteraciones para la misma precisión"),
]
x0, y0 = 0.75, 1.55
widths = [2.2, 4.65, 4.95]
for j, h in enumerate(headers):
    box(s, x0 + sum(widths[:j]), y0, widths[j], 0.7, PURPLE)
    text(s, h, x0 + sum(widths[:j]) + 0.05, y0 + 0.06, widths[j] - 0.1, 0.55, 15, WHITE, True, PP_ALIGN.CENTER)
for i, row in enumerate(data):
    for j, value in enumerate(row):
        box(s, x0 + sum(widths[:j]), y0 + 0.7 + i * 0.75, widths[j], 0.75, PANEL, False, RGBColor(45, 55, 78))
        text(s, value, x0 + sum(widths[:j]) + 0.08, y0 + 0.76 + i * 0.75, widths[j] - 0.16, 0.6, 14, CYAN if j == 0 else WHITE, j == 0)
text(s, "No existe un método universalmente mejor: la elección depende del sistema y de la información disponible.", 0.95, 6.25, 11.4, 0.4, 16, ORANGE, True, PP_ALIGN.CENTER)

# 10
s = base(prs, "Conclusiones y cierre", "Resumen", 10)
box(s, 0.8, 1.35, 7.2, 4.9, PANEL, True)
text(s, "Ideas clave para recordar", 1.15, 1.7, 6.4, 0.4, 22, CYAN, True)
bullet(s, "Ambos métodos buscan F(x)=0 mediante aproximaciones sucesivas.", 1.15, 2.4, 6.3, size=17)
bullet(s, "Newton usa el Jacobiano y una corrección lineal.", 1.15, 3.13, 6.3, size=17)
bullet(s, "Gauss-Seidel usa despejes y valores actualizados al instante.", 1.15, 3.86, 6.3, size=17)
bullet(s, "Siempre se debe revisar error, residuo y convergencia.", 1.15, 4.59, 6.3, size=17)
box(s, 8.45, 1.35, 4.05, 4.9, PANEL, True)
text(s, "Mensaje final", 8.8, 1.78, 3.35, 0.35, 21, PURPLE, True, PP_ALIGN.CENTER)
text(s, "La solución numérica\nno es solo calcular:\n\ntambién es elegir una\nbuena formulación,\nuna semilla adecuada\ny un criterio de parada.", 8.82, 2.42, 3.3, 2.65, 19, WHITE, True, PP_ALIGN.CENTER)
text(s, "¿Preguntas?", 1.0, 6.45, 11.4, 0.4, 24, ORANGE, True, PP_ALIGN.CENTER)

prs.save(OUT)
print(OUT)
