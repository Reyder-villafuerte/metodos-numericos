from io import BytesIO

from openpyxl import load_workbook
import pytest

from metodos.muller import muller
from utilidades.exportacion import crear_excel


def funcion_clase(x):
    return x ** 3 - 13 * x - 12


def test_muller_ejemplo_de_clase():
    resultado = muller(funcion_clase, 4.5, 5.5, 5.0,
                       error_permitido=0.5, max_iteraciones=100, criterio="error")
    primera = resultado.iteraciones[0]
    assert primera["a"] == pytest.approx(15.0)
    assert primera["b"] == pytest.approx(62.25)
    assert primera["c"] == pytest.approx(48.0)
    assert primera["x3"] == pytest.approx(3.9765, abs=1e-4)
    assert resultado.iteraciones[1]["x3"] == pytest.approx(4.001, abs=1e-3)
    assert resultado.raiz == pytest.approx(4.0, abs=1e-4)
    assert resultado.error_final < 0.5
    assert resultado.convergio is True


def test_muller_iteraciones_fijas():
    resultado = muller(funcion_clase, 4.5, 5.5, 5.0, criterio="fijo",
                       iteraciones_fijas=2, error_permitido=None, max_iteraciones=None)
    assert len(resultado.iteraciones) == 2
    assert resultado.convergio is None
    assert resultado.raiz == pytest.approx(4.001, abs=1e-3)


def test_muller_encuentra_raiz_compleja():
    resultado = muller(lambda x: x ** 2 + 1, 0, 1, 2,
                       error_permitido=0.01, max_iteraciones=20, criterio="error")
    raiz = complex(resultado.raiz)
    assert abs(raiz.real) < 1e-10
    assert abs(abs(raiz.imag) - 1) < 1e-10
    assert resultado.convergio is True


def test_muller_rechaza_valores_iniciales_repetidos():
    with pytest.raises(ValueError, match="deben ser distintos"):
        muller(funcion_clase, 1, 1, 2, error_permitido=1, max_iteraciones=10)


def test_muller_detecta_denominadores_finales_cero():
    with pytest.raises(ValueError, match="dos posibles denominadores"):
        muller(lambda x: 1, 0, 1, 2, error_permitido=1, max_iteraciones=10)


def test_muller_detecta_denominador_de_coeficientes_casi_cero():
    with pytest.raises(ValueError, match="calcular a y b"):
        muller(funcion_clase, 0, 1e-7, 2e-7, error_permitido=1, max_iteraciones=10)


def test_muller_reporta_maximo_de_iteraciones():
    resultado = muller(funcion_clase, 4.5, 5.5, 5.0,
                       error_permitido=1e-30, max_iteraciones=1, criterio="error")
    assert resultado.convergio is False
    assert len(resultado.iteraciones) == 1
    assert "máximo" in resultado.motivo


def _exportar(resultado, funcion):
    contenido = crear_excel(
        "Prueba Müller", funcion, "Müller", {"x0": 0, "x1": 1, "x2": 2},
        0.01, resultado.raiz, resultado.error_final, resultado.iteraciones,
        "Error porcentual permitido", "Convergió",
    )
    return load_workbook(BytesIO(contenido), data_only=True)


def test_exportacion_muller_real():
    resultado = muller(funcion_clase, 4.5, 5.5, 5.0,
                       error_permitido=0.5, max_iteraciones=100)
    libro = _exportar(resultado, "x^3-13*x-12")
    assert {"Resumen", "Iteraciones", "Descripción"}.issubset(libro.sheetnames)
    assert libro["Iteraciones"].max_row == len(resultado.iteraciones) + 1


def test_exportacion_muller_complejo_usa_texto_legible():
    resultado = muller(lambda x: x ** 2 + 1, 0, 1, 2,
                       error_permitido=0.01, max_iteraciones=20)
    libro = _exportar(resultado, "x^2+1")
    valores = [celda.value for fila in libro["Resumen"].iter_rows() for celda in fila]
    assert any(isinstance(valor, str) and "i" in valor for valor in valores)
