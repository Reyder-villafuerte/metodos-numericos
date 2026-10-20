from io import BytesIO
import logging
import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment
import math
import sympy as sp


LOGGER = logging.getLogger("metodos_numericos.exportacion")


def _convertir_valor(valor):
    """Convierte un valor a un tipo Python nativo seguro para Excel."""
    if valor is None:
        return ""
    if isinstance(valor, (int, float)):
        if math.isnan(valor) or math.isinf(valor):
            return "Error"
        return valor
    if isinstance(valor, sp.Basic):
        return str(valor)
    if isinstance(valor, complex):
        signo = "+" if valor.imag >= 0 else "-"
        return f"{valor.real:.10f} {signo} {abs(valor.imag):.10f}i"
    return valor


def crear_excel(titulo, funcion, metodo, parametros, error_permitido, raiz, error_final, iteraciones,
                criterio=None, estado=None):
    """Crea un archivo Excel con resultados de métodos numéricos."""
    salida = BytesIO()
    
    try:
        # Preparar datos del resumen
        raiz_conv = _convertir_valor(raiz) if raiz is not None else 0.0
        error_final_conv = float(error_final) if error_final is not None else 0.0
        error_permitido_conv = float(error_permitido) if error_permitido is not None else 0.0
        
        # Crear DataFrame del resumen
        nombre_valor_criterio = "Iteraciones solicitadas" if criterio and "fijo" in criterio.lower() else "Error permitido (%)"
        resumen_data = [
            {"Descripción": "Título", "Valor": titulo},
            {"Descripción": "Función original", "Valor": str(funcion)},
            {"Descripción": "Método", "Valor": metodo},
            {"Descripción": "Criterio de detención", "Valor": criterio or "Error permitido"},
            {"Descripción": nombre_valor_criterio, "Valor": error_permitido_conv},
            {"Descripción": "Raíz aproximada", "Valor": raiz_conv},
            {"Descripción": "Error final (%)", "Valor": error_final_conv},
            {"Descripción": "Estado", "Valor": estado or "No especificado"},
            {"Descripción": "Número de iteraciones", "Valor": len(iteraciones) if iteraciones else 0},
        ]
        if isinstance(parametros, dict):
            filas_parametros = [
                {"Descripción": clave, "Valor": _convertir_valor(valor)}
                for clave, valor in parametros.items()
                if clave in {"a", "b", "x0", "x1", "x2"}
            ]
            resumen_data[3:3] = filas_parametros
        else:
            resumen_data.insert(3, {"Descripción": "Parámetros iniciales", "Valor": str(parametros)})
        
        # Preparar datos de iteraciones
        iteraciones_conv = []
        if iteraciones:
            for fila in iteraciones:
                fila_conv = {}
                for clave, valor in fila.items():
                    fila_conv[str(clave)] = _convertir_valor(valor)
                iteraciones_conv.append(fila_conv)
        
        # Crear archivo Excel
        with pd.ExcelWriter(salida, engine="openpyxl") as escritor:
            # Hoja de resumen
            df_resumen = pd.DataFrame(resumen_data)
            df_resumen.to_excel(escritor, sheet_name="Resumen", index=False)
            
            # Hoja de iteraciones
            if iteraciones_conv:
                df_iteraciones = pd.DataFrame(iteraciones_conv)
                df_iteraciones.to_excel(escritor, sheet_name="Iteraciones", index=False)
            
            # Hoja de descripción
            descripcion_data = [
                {"Sección": "Método", "Detalle": metodo},
                {"Sección": "Función", "Detalle": str(funcion)},
                {"Sección": "Criterio de parada", "Detalle": f"Error permitido: {error_permitido_conv}%"},
                {"Sección": "Raíz encontrada", "Detalle": str(raiz_conv)},
                {"Sección": "Error final", "Detalle": f"{error_final_conv}%"},
                {"Sección": "Interpretación", "Detalle": "La raíz es el valor de x donde f(x) = 0 o está cercano a él."},
            ]
            df_descripcion = pd.DataFrame(descripcion_data)
            df_descripcion.to_excel(escritor, sheet_name="Descripción", index=False)
            
            # Formatear hojas
            for sheet_name in escritor.sheets:
                ws = escritor.sheets[sheet_name]
                for row in ws.iter_rows(min_row=1, max_row=ws.max_row, min_col=1, max_col=ws.max_column):
                    for cell in row:
                        if cell.row == 1:
                            cell.font = Font(bold=True, color="FFFFFF")
                            cell.fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
                        cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
                
                # Ajustar ancho de columnas
                for col in ws.columns:
                    max_length = 0
                    col_letter = col[0].column_letter
                    for cell in col:
                        if len(str(cell.value or "")) > max_length:
                            max_length = len(str(cell.value or ""))
                    ws.column_dimensions[col_letter].width = min(max_length + 2, 50)
        
        return salida.getvalue()
    
    except Exception as exc:
        LOGGER.exception("Error al crear el archivo Excel: %s", exc)
        raise ValueError(f"Error al preparar el archivo Excel: {str(exc)}") from exc
